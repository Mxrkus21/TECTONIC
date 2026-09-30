"""Convert downloaded document bytes into normalized plain text."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path
from typing import Callable, Dict, Optional
from xml.etree import ElementTree

from contracts import ExtractedContent


DEFAULT_MAX_CHARACTERS = 250_000


class ExtractionError(Exception):
    """Base error for content that could not be converted to text."""


class UnsupportedFileTypeError(ExtractionError):
    pass


class OcrRequiredError(ExtractionError):
    pass


class MissingExtractorDependencyError(ExtractionError):
    pass


def extract_content(
    filename: str,
    content: bytes,
    mime_type: Optional[str] = None,
    max_characters: int = DEFAULT_MAX_CHARACTERS,
) -> ExtractedContent:
    """Extract and bound document text without inventing missing content."""

    if max_characters < 1:
        raise ValueError("max_characters must be greater than zero")
    if not content:
        raise ExtractionError(f"Document is empty: {filename}")

    extension = _resolve_extension(filename, mime_type)
    extractor = _extractors().get(extension)
    if extractor is None:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{extension or 'unknown'}' for {filename}"
        )

    try:
        text = extractor(content)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Failed to extract {filename}: {exc}") from exc

    normalized = normalize_text(text)
    if not normalized:
        if extension == ".pdf":
            raise OcrRequiredError(f"PDF contains no extractable text: {filename}")
        raise ExtractionError(f"No extractable text found in {filename}")

    character_count = len(normalized)
    truncated = character_count > max_characters
    visible_text = normalized[:max_characters] if truncated else normalized
    return ExtractedContent(
        full_text=visible_text,
        relevant_passages=[],
        character_count=character_count,
        truncated=truncated,
        extraction_method="local",
    )


def _extractors() -> Dict[str, Callable[[bytes], str]]:
    return {
        ".txt": _decode_text,
        ".md": _decode_text,
        ".csv": _extract_csv,
        ".json": _extract_json,
        ".docx": _extract_docx,
        ".pdf": _extract_pdf,
        ".xlsx": _extract_xlsx,
    }


def _resolve_extension(filename: str, mime_type: Optional[str]) -> str:
    extension = Path(filename).suffix.lower()
    if extension:
        return extension

    mime_extensions = {
        "text/plain": ".txt",
        "text/markdown": ".md",
        "text/csv": ".csv",
        "application/json": ".json",
        "application/pdf": ".pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    }
    return mime_extensions.get((mime_type or "").lower(), "")


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Text encoding could not be detected")


def _extract_csv(content: bytes) -> str:
    decoded = _decode_text(content)
    rows = list(csv.reader(io.StringIO(decoded)))
    return "\n".join(",".join(cell.strip() for cell in row) for row in rows)


def _extract_json(content: bytes) -> str:
    decoded = _decode_text(content)
    try:
        value = json.loads(decoded)
    except json.JSONDecodeError as exc:
        raise ExtractionError(f"Invalid JSON: {exc.msg}") from exc
    return json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True)


def _extract_docx(content: bytes) -> str:
    paragraphs = []
    word_namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            xml_files = ["word/document.xml"]
            xml_files.extend(
                name
                for name in archive.namelist()
                if name.startswith(("word/header", "word/footer"))
                and name.endswith(".xml")
            )
            for xml_name in xml_files:
                if xml_name not in archive.namelist():
                    continue
                root = ElementTree.fromstring(archive.read(xml_name))
                for paragraph in root.iter(f"{word_namespace}p"):
                    text = "".join(
                        node.text or ""
                        for node in paragraph.iter(f"{word_namespace}t")
                    ).strip()
                    if text:
                        paragraphs.append(text)
    except (zipfile.BadZipFile, ElementTree.ParseError, KeyError) as exc:
        raise ExtractionError("Invalid DOCX document") from exc

    return "\n".join(paragraphs)


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise MissingExtractorDependencyError(
            "PDF extraction requires the 'pypdf' package"
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ExtractionError("Invalid or unreadable PDF document") from exc
    return "\n\n".join(pages)


def _extract_xlsx(content: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise MissingExtractorDependencyError(
            "XLSX extraction requires the 'openpyxl' package"
        ) from exc

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        output = []
        for worksheet in workbook.worksheets:
            output.append(f"## Sheet: {worksheet.title}")
            for row in worksheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    output.append(",".join(values))
        workbook.close()
        return "\n".join(output)
    except Exception as exc:
        raise ExtractionError("Invalid or unreadable XLSX document") from exc


def normalize_text(value: str) -> str:
    """Normalize extracted text consistently across local and cloud extractors."""

    lines = [" ".join(line.split()) for line in value.replace("\x00", "").splitlines()]
    output = []
    previous_blank = False
    for line in lines:
        is_blank = not line
        if is_blank and previous_blank:
            continue
        output.append(line)
        previous_blank = is_blank
    return "\n".join(output).strip()
